"""Render a tailored CV (dict) as a clean, ATS-friendly .docx."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def _rule(paragraph):
    """Thin bottom border under a section heading."""
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3A5F")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)


def _section(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title.upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(11), ACCENT
    _rule(p)


def _bullets(doc, items):
    for b in items or []:
        if b:
            p = doc.add_paragraph(str(b), style="List Bullet")
            p.paragraph_format.space_after = Pt(1)


def _line(doc, bold, rest="", space_before=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(1)
    p.add_run(bold).bold = True
    if rest:
        r = p.add_run(f"  |  {rest}")
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return p


def render_docx(cv: dict, path: Path) -> Path:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.6)
        s.left_margin = s.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(cv.get("name", ""))
    r.bold, r.font.size, r.font.color.rgb = True, Pt(20), ACCENT
    if cv.get("headline"):
        p = doc.add_paragraph(cv["headline"])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.size = Pt(11.5)
    contact = [str(v) for v in (cv.get("contact") or {}).values() if v]
    if contact:
        p = doc.add_paragraph("  |  ".join(contact))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.size = Pt(9.5)

    if cv.get("summary"):
        _section(doc, "Professional Summary")
        doc.add_paragraph(cv["summary"])

    if cv.get("skills"):
        _section(doc, "Technical Skills")
        for g in cv["skills"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.add_run(f"{g.get('category', '')}: ").bold = True
            p.add_run(", ".join(g.get("items") or []))

    if cv.get("experience"):
        _section(doc, "Experience")
        for e in cv["experience"]:
            head = " — ".join(x for x in (e.get("title"), e.get("company")) if x)
            meta = "  |  ".join(x for x in (e.get("location"), e.get("dates")) if x)
            _line(doc, head, meta)
            _bullets(doc, e.get("bullets"))

    if cv.get("projects"):
        _section(doc, "Projects")
        for pr in cv["projects"]:
            meta = ", ".join(pr.get("tech") or [])
            _line(doc, pr.get("name", ""), meta)
            if pr.get("link"):
                doc.add_paragraph(pr["link"]).runs[0].font.size = Pt(9)
            _bullets(doc, pr.get("bullets"))

    if cv.get("education"):
        _section(doc, "Education")
        for ed in cv["education"]:
            head = " — ".join(x for x in (ed.get("degree"), ed.get("institution")) if x)
            _line(doc, head, ed.get("dates", ""))
            if ed.get("details"):
                doc.add_paragraph(ed["details"])

    for title, key in (("Certifications", "certifications"), ("Languages", "languages")):
        if cv.get(key):
            _section(doc, title)
            _bullets(doc, cv[key])

    for extra in cv.get("extra_sections") or []:
        if extra.get("items"):
            _section(doc, extra.get("title", ""))
            _bullets(doc, extra["items"])

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
