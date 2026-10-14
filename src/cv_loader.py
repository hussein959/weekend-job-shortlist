"""Load the master CV from a private share link (CV_URL secret).

The CV is never stored in this repository. CV_URL should be a Google Drive file
link (PDF or .docx) or a Google Docs link, shared as "Anyone with the link can view".
Dropbox and direct download links also work.

Supported formats: .pdf, .docx, Google Docs, .md, .txt.

For local runs only, CV_PATH may point to a file on your own machine instead.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

import requests

from .common import env, log


class CVError(SystemExit):
    pass


def _direct_link(url: str) -> str:
    m = re.search(r"drive\.google\.com/(?:file/d/|open\?id=|uc\?(?:.*&)?id=)([\w-]+)", url)
    if m:
        return f"https://drive.google.com/uc?export=download&id={m.group(1)}"
    m = re.search(r"docs\.google\.com/document/d/([\w-]+)", url)
    if m:
        return f"https://docs.google.com/document/d/{m.group(1)}/export?format=docx"
    if "dropbox.com" in url:
        return re.sub(r"([?&])dl=0", r"\1dl=1", url) if "dl=0" in url else url + ("&" if "?" in url else "?") + "dl=1"
    return url


def _extract(data: bytes, name: str) -> str:
    name = name.lower()
    if data[:4] == b"%PDF" or name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
    if data[:2] == b"PK" or name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return data.decode("utf-8", errors="replace")


def _download(url: str) -> bytes:
    # Errors are reported without the URL: logs of a public repo are public, and
    # GitHub only masks the secret's exact value, not the converted download link.
    try:
        r = requests.get(_direct_link(url), timeout=60, allow_redirects=True)
    except requests.RequestException as e:
        raise CVError(f"Could not download the CV from CV_URL ({type(e).__name__}).") from None
    if r.status_code != 200:
        raise CVError(f"Downloading the CV from CV_URL failed with HTTP {r.status_code}. "
                      "Check the link and that it is shared as 'Anyone with the link can view'.")
    head = r.content[:200].lstrip().lower()
    if head.startswith(b"<!doctype html") or head.startswith(b"<html"):
        raise CVError("CV_URL returned a web page instead of the file. Make sure the link is shared "
                      "as 'Anyone with the link can view' and points to a single file or Google Doc.")
    return r.content


def load_cv() -> str:
    url = env("CV_URL")
    local = env("CV_PATH")
    if url:
        text = _extract(_download(url), "")
    elif local and not env("GITHUB_ACTIONS"):
        path = Path(local).expanduser()
        if not path.exists():
            raise CVError(f"CV_PATH does not exist: {path}")
        text = _extract(path.read_bytes(), path.name)
    else:
        raise CVError("Missing CV_URL secret: add your Google Drive / Google Docs share link "
                      "as a repository secret named CV_URL.")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 300:
        raise CVError("CV text is suspiciously short; check that CV_URL points to your CV "
                      "(a text-based PDF, .docx or Google Doc, not a scanned image).")
    log.info("Loaded master CV (%d characters)", len(text))
    return text
